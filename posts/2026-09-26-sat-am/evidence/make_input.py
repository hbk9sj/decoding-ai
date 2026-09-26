# Builds one small regional sales table and saves it as PDF, DOCX, XLSX and PPTX.
import csv
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, Paragraph, Spacer, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import docx, openpyxl
from pptx import Presentation
from pptx.util import Inches

HEAD = ["Region", "Q1 revenue", "Q2 revenue", "Change", "Owner"]
ROWS = [
    ["North", "41,200", "38,900", "-5.6%", "A. Mehta"],
    ["South", "52,750", "61,300", "16.2%", "R. Iyer"],
    ["East", "29,400", "29,950", "1.9%", "S. Das"],
    ["West", "47,100", "44,020", "-6.5%", "K. Shah"],
    ["Central", "18,600", "23,480", "26.2%", "P. Rao"],
    ["Online", "66,300", "71,150", "7.3%", "D. Nair"],
]
TITLE = "Regional sales, Q1 vs Q2 2026 (INR thousands)"
with open("table.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(HEAD); w.writerows(ROWS)

# PDF
st = getSampleStyleSheet()
doc = SimpleDocTemplate("report.pdf", pagesize=A4)
t = Table([HEAD] + ROWS)
t.setStyle(TableStyle([("GRID", (0,0), (-1,-1), 0.5, colors.grey), ("BACKGROUND", (0,0), (-1,0), colors.lightgrey)]))
doc.build([Paragraph(TITLE, st["Heading2"]), Paragraph("Prepared for the Q2 review.", st["Normal"]), Spacer(1, 12), t])

# DOCX
d = docx.Document(); d.add_heading(TITLE, 2); d.add_paragraph("Prepared for the Q2 review.")
tb = d.add_table(rows=1+len(ROWS), cols=len(HEAD))
for i, r in enumerate([HEAD] + ROWS):
    for j, v in enumerate(r): tb.cell(i, j).text = v
d.save("report.docx")

# XLSX
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q2"
ws.append(HEAD)
for r in ROWS: ws.append(r)
wb.save("report.xlsx")

# PPTX
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = TITLE
gt = s.shapes.add_table(1+len(ROWS), len(HEAD), Inches(0.5), Inches(1.5), Inches(9), Inches(4)).table
for i, r in enumerate([HEAD] + ROWS):
    for j, v in enumerate(r): gt.cell(i, j).text = v
p.save("report.pptx")
print("built report.pdf report.docx report.xlsx report.pptx table.csv")
